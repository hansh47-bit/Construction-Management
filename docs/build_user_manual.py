from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name('INC_공사통합관리_사용자매뉴얼_v1.0.docx')

BLUE = '1F4E79'
MID_BLUE = 'D9EAF7'
PALE_BLUE = 'F4F8FC'
PALE_GRAY = 'F5F6F8'
GRID = 'D9D9D9'
TEXT = '202124'
MUTED = '5B6573'
GREEN = 'E8F5ED'
AMBER = 'FFF4D6'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=GRID, size='6'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_width(cell, inches):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(inches * 1440)))
    tcW.set(qn('w:type'), 'dxa')

def style_run(run, size=10.5, bold=False, color=TEXT):
    run.font.name = 'Malgun Gothic'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_text(cell, text, size=9.5, bold=False, color=TEXT, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    style_run(run, size, bold, color)
    return p

def add_para(doc, text='', size=10.5, bold=False, color=TEXT, space_after=6, align=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    style_run(r, size, bold, color)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        style_run(p.add_run(item), 10.2, False)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    style_run(r, 16 if level == 1 else 12.5, True, '000000')
    return p

def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths: set_width(cell, widths[i])
        add_text(cell, text, font_size, True, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, PALE_BLUE if r_idx % 2 else 'FFFFFF')
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: set_width(cell, widths[i])
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) <= 4 else None
            add_text(cell, str(text), font_size, False, TEXT, align)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def screen_frame(doc, title, subtitle, nav_items, cards=None, columns=None, note=None):
    # A compact native Word representation of the actual screen structure.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run('화면 예시  ' + title)
    style_run(r, 10.5, True, '000000')
    outer = doc.add_table(rows=1, cols=1)
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer.autofit = False
    cell = outer.cell(0, 0)
    set_width(cell, 6.85)
    set_cell_border(cell, 'AAB7C4', '10')
    set_cell_margins(cell, 0, 0, 0, 0)
    set_cell_shading(cell, 'FFFFFF')
    header = cell.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.space_before = Pt(0)
    r = header.add_run('  아이엔씨 ERP   |   공사 통합관리')
    style_run(r, 9.5, True, 'FFFFFF')
    pPr = header._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), BLUE); pPr.append(shd)
    nav = cell.add_table(rows=1, cols=len(nav_items))
    nav.autofit = False
    for i, item in enumerate(nav_items):
        c = nav.cell(0, i)
        set_cell_border(c, 'C7D5E2', '4'); set_cell_margins(c, 70, 70, 70, 70)
        set_cell_shading(c, MID_BLUE if i == 0 else 'FFFFFF')
        add_text(c, item, 8.2, i == 0, BLUE if i == 0 else MUTED, WD_ALIGN_PARAGRAPH.CENTER)
    body = cell.add_paragraph()
    body.paragraph_format.space_before = Pt(7); body.paragraph_format.space_after = Pt(4)
    style_run(body.add_run('  ' + subtitle), 11.5, True, TEXT)
    if cards:
        grid = cell.add_table(rows=1, cols=len(cards))
        grid.autofit = False
        for i, (label, value, helper) in enumerate(cards):
            c = grid.cell(0, i)
            set_cell_border(c); set_cell_margins(c, 100, 100, 100, 100); set_cell_shading(c, PALE_BLUE)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            style_run(p.add_run(label), 7.8, True, MUTED)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            style_run(p.add_run(value), 10, True, BLUE)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0)
            style_run(p.add_run(helper), 7.2, False, MUTED)
    if columns:
        grid = cell.add_table(rows=1, cols=len(columns))
        grid.autofit = False
        for i, (head, lines) in enumerate(columns):
            c = grid.cell(0, i)
            set_cell_border(c); set_cell_margins(c, 90, 100, 90, 100); set_cell_shading(c, 'FFFFFF')
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
            style_run(p.add_run(head), 8.4, True, BLUE)
            for line in lines:
                p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                style_run(p.add_run(line), 7.8, False, TEXT)
    if note:
        p = cell.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)
        style_run(p.add_run('  ' + note), 7.8, False, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run('아이엔씨 공사 통합관리 시스템 사용자 매뉴얼  |  내부 업무용')
    style_run(r, 8, False, '6B7280')

def set_core_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Malgun Gothic'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    for s in ('Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[s]
        style.font.name = 'Malgun Gothic'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
        style.font.color.rgb = RGBColor(0, 0, 0)

def new_page(doc):
    doc.add_page_break()

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.58)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
set_core_styles(doc)
add_footer(section)

# Cover
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(122)
p.paragraph_format.space_after = Pt(14)
r = p.add_run('아이엔씨 공사 통합관리 시스템 사용자 매뉴얼')
style_run(r, 25, True, '000000')
p = add_para(doc, '프로젝트 등록부터 준공정산까지', 14, False, MUTED, 8, WD_ALIGN_PARAGRAPH.CENTER)
p = add_para(doc, '계약  실행예산  발주  기성  지급  준공정산 업무 흐름 안내', 11, False, MUTED, 40, WD_ALIGN_PARAGRAPH.CENTER)
meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (label, value) in enumerate([('문서 버전', 'v1.0'), ('기준일', '2026년 10월 28일'), ('대상 사용자', '대표이사  팀장  현장담당자  관리 회계 담당자')]):
    for j, text in enumerate((label, value)):
        c = meta.cell(i,j); set_cell_border(c); set_cell_margins(c, 120, 140, 120, 140)
        set_cell_shading(c, MID_BLUE if j == 0 else 'FFFFFF')
        add_text(c, text, 10, j == 0, TEXT, WD_ALIGN_PARAGRAPH.CENTER if j == 0 else None)

new_page(doc)
add_heading(doc, '1  이 매뉴얼을 사용하는 방법')
add_para(doc, '이 문서는 실제 시스템에 표시되는 메뉴명과 입력 흐름을 기준으로 작성했습니다. 먼저 프로젝트를 선택한 뒤, 상단 통합현황과 프로젝트별 탭을 따라 업무를 진행하면 됩니다.', 10.5)
add_table(doc, ['표기', '의미', '사용 방법'], [
    ['화면 예시', '실제 UI의 메뉴와 주요 정보 배치를 단순화해 보여주는 구성도', '화면에서 찾을 위치를 먼저 확인합니다.'],
    ['입력', '사용자가 직접 작성하거나 선택하는 필드', '금액, 사유, 담당자, 첨부파일을 빠짐없이 입력합니다.'],
    ['결재요청', '대표 또는 결재권자 검토가 필요한 단계', '저장 후 통합결재함에서 상태를 확인합니다.'],
    ['주의', '금액 흐름과 이력 보존을 위한 업무 규칙', '승인본을 덮어쓰지 않고 변경 차수 또는 변경안을 사용합니다.'],
], [1.1, 2.55, 3.2])
add_heading(doc, '전체 업무 흐름', 2)
add_table(doc, ['순서', '업무 단계', '주요 산출물', '확인 화면'], [
    ['1', '프로젝트 등록', '프로젝트 번호와 기본 일정', '프로젝트 현황'],
    ['2', '고객 계약', '계약 공종과 계약금액', '고객 견적 계약 관리'],
    ['3', '세부실행예산', '공종별 실행항목과 원가', '세부실행예산'],
    ['4', '업체 비교와 발주', '선정업체와 발주품의', '업체 비교견적'],
    ['5', '기성과 지급', '차수별 기성 및 지급결재', '기성 및 지급 현황'],
    ['6', '준공정산', '최종 원가와 이익 및 결재', '준공정산'],
], [0.55, 1.45, 2.3, 2.55])

new_page(doc)
add_heading(doc, '2  처음 시작하기')
add_heading(doc, '2 1  상단 메뉴와 프로젝트 선택', 2)
screen_frame(doc, '프로젝트 목록', '프로젝트 현황', ['프로젝트 현황', '통합결재함', '변경관리 이력'],
    columns=[('프로젝트명 프로젝트번호', ['성가수녀원 리뉴얼공사', 'P-2026-001']), ('고객 현장', ['성가수녀원', '서울시 성북구']), ('일정 상태', ['2026-03-10 시작', '공사진행'])],
    note='프로젝트명 또는 프로젝트번호로 검색한 뒤, 목록의 프로젝트 행을 선택하면 통합현황 상세로 이동합니다.')
add_bullets(doc, [
    '프로젝트 현황에서 프로젝트 등록을 누르면 신규 프로젝트의 기본정보를 등록할 수 있습니다.',
    '프로젝트 목록의 프로젝트명 행을 누르면 해당 프로젝트의 통합현황과 세부 업무 탭이 열립니다.',
    '통합결재함은 여러 프로젝트의 결재를 한 화면에서 확인하는 전역 메뉴입니다.',
    '변경관리 이력은 먼저 프로젝트를 선택해야 해당 프로젝트의 차수를 조회하고 작성할 수 있습니다.',
])
add_heading(doc, '2 2  프로젝트 등록 시 필수 확인 항목', 2)
add_table(doc, ['항목', '입력 기준', '업무상 의미'], [
    ['프로젝트명과 고객 발주처', '계약서의 공식 명칭을 사용', '목록과 결재문서의 식별 기준'],
    ['계약일 착공일 준공예정일', '실제 계약 및 현장 일정 기준', '프로젝트 진행상태 판단'],
    ['담당 팀장과 담당자', '내부 담당자를 지정', '현장 기록과 책임자 확인'],
    ['목표 이익률', '계약 검토 시 합의된 비율', '목표원가한도를 자동 산정'],
], [1.5, 2.6, 2.75])

new_page(doc)
add_heading(doc, '3  프로젝트 통합현황 읽기')
add_para(doc, '프로젝트를 선택하면 계약부터 지급까지의 현재 승인 기준 금액을 한 화면에서 확인할 수 있습니다. 상단의 수주 및 목표 현황과 실제 자금 집행 현황을 함께 읽는 것이 핵심입니다.')
screen_frame(doc, '통합현황', '성가수녀원 리뉴얼공사  공사진행', ['고객 견적 계약 관리', '세부실행예산', '업체 비교견적', '기성 및 지급 현황', '준공정산'],
    cards=[('현재 계약금액', '220,000,000 원', '현재 승인 계약 공급가액'), ('목표 원가한도', '121,000,000 원', '목표 이익률 기준'), ('예상 이익 이익률', '103,780,000 원  47.2%', '계약금액 - 예상 최종원가')],
    columns=[('실행과 발주', ['현재 실행예산 6,020,000 원', '현재 발주금액 110,200,000 원']), ('기성과 지급', ['누적 기성금액 80,000,000 원', '실제 지급금액 70,000,000 원']), ('잔액', ['총 미지급 40,200,000 원', '승인 후 미지급 10,000,000 원'])],
    note='모든 집계는 현재 승인 차수를 기준으로 표시됩니다. 이전 차수의 기록은 보존되지만 현재 총액에는 포함되지 않습니다.')
add_table(doc, ['지표', '확인 방법', '주의점'], [
    ['목표원가한도', '계약 공급가액과 목표 이익률로 자동 계산', '계약 변경 시 함께 검토합니다.'],
    ['예상 최종원가', '확정 발주와 미발주 실행예산 및 직접 정산비 반영', '예상 이익을 판단하는 기준입니다.'],
    ['미기성금액', '현재 발주금액에서 승인 기성누계를 뺀 금액', '아직 승인되지 않은 공정 잔액입니다.'],
    ['승인 후 미지급', '승인 기성누계에서 실제 지급누계를 뺀 금액', '지급결재가 필요한 금액입니다.'],
], [1.25, 3.05, 2.55])

new_page(doc)
add_heading(doc, '4  고객 견적과 계약 관리')
screen_frame(doc, '고객 견적 계약 관리', '공급가액 기준 계약금액과 목표이익률', ['고객 견적 계약 관리', '세부실행예산', '업체 비교견적'],
    cards=[('계약 총 공급가액', '220,000,000 원', '최초 계약 및 승인 변경 반영'), ('목표 이익', '99,000,000 원', '목표 이익률 45 퍼센트'), ('목표원가한도', '121,000,000 원', '계약 공급가액 기준')],
    columns=[('계약 공종표', ['공종', '견적 제출 공급가액', '최종 계약 공급가액']), ('자동 계산', ['VAT 10 퍼센트', 'VAT 포함 총액', '합계']), ('증빙', ['견적서 첨부', '계약서 첨부', '계약 변경안 작성'])],
    note='계약 공급가액을 입력하면 VAT가 자동 계산됩니다. 승인된 최초 금액은 수정하지 말고 계약 변경안을 작성합니다.')
add_heading(doc, '처리 순서', 2)
add_bullets(doc, [
    '공종별 견적 제출 금액과 최종 계약 공급가액을 입력합니다.',
    '목표 이익률을 확인해 목표원가한도가 현실적인지 검토합니다.',
    '견적서와 계약서를 구분해 첨부한 뒤 계약 저장을 누릅니다.',
    '계약금액 또는 계약 공종을 바꿔야 하면 계약 변경안 작성을 사용하고, 변경관리 이력에서 차수를 관리합니다.',
])

new_page(doc)
add_heading(doc, '5  세부실행예산 작성과 결재요청')
screen_frame(doc, '세부실행예산', '최초 승인본 세부실행예산 항목', ['세부실행예산', '업체 비교견적', '계약 발주 이력'],
    cards=[('계약 공급가액', '220,000,000 원', '고객 계약 기준'), ('목표원가한도', '121,000,000 원', '목표이익률 기준'), ('세부실행예산 총합계', '6,020,000 원', '현재 작성 차수')],
    columns=[('실행항목ID별 집계', ['EX-001', '인테리어공사', '외주비 폐기물 인건비']), ('세부 항목', ['원가구분', '규격 단위 수량 단가', '금액과 비고']), ('완료 작업', ['행 추가 또는 행 복사', '세부실행예산 저장', '결재요청'])],
    note='발주가 필요한 항목은 승인된 실행예산에서만 업체 비교견적의 선택 대상으로 표시됩니다.')
add_table(doc, ['입력 항목', '작성 기준', '자동 반영'], [
    ['실행항목ID', '같은 발주 단위는 같은 ID로 입력', 'ID별 금액이 자동 집계됩니다.'],
    ['원가구분', '외주비 폐기물 인건비 등 실제 성격 선택', '공종별 원가분석에 반영됩니다.'],
    ['수량과 단가', '부가세 포함 여부를 혼동하지 않도록 기준 통일', '금액이 자동 계산됩니다.'],
    ['결재 첨부파일', '산출근거와 견적자료를 첨부', '결재권자가 검토할 근거가 됩니다.'],
], [1.35, 3.0, 2.5])
add_para(doc, '주의  실행예산을 저장했다고 발주할 수 있는 것은 아닙니다. 결재요청 후 승인된 실행예산 항목만 업체 비교견적 화면에서 선택됩니다.', 10, True, BLUE, 5)

new_page(doc)
add_heading(doc, '6  업체 비교견적과 발주품의')
screen_frame(doc, '업체 비교견적', '승인 실행예산 기준 업체 비교와 발주품의', ['업체 비교견적', '계약 발주 이력', '기성 및 지급 현황'],
    cards=[('실행항목ID', 'EX-001', '승인된 실행예산 항목'), ('승인 실행예산', '6,020,000 원', '롤업 금액 기준'), ('발주품의 상태', '미요청', '선정 후 결재요청')],
    columns=[('비교견적 입력', ['업체명', '견적금액', '공사기간과 결제조건']), ('선정', ['한 업체 선택', '선정사유 입력', '견적서 파일 첨부']), ('발주 연결', ['비교견적 저장', '발주품의 요청', '통합결재 승인 후 발주 확정'])],
    note='발주품의 요청 후에는 비교견적을 임의로 바꿀 수 없습니다. 선정되지 않은 업체는 비교 이력으로 남고 발주 및 지급에는 연결되지 않습니다.')
add_heading(doc, '처리 순서', 2)
add_bullets(doc, [
    '실행항목ID 선택 목록에서 승인된 실행예산 항목을 고릅니다.',
    '업체별 견적금액, 공사기간, 결제조건을 입력하고 비교견적을 저장합니다.',
    '선정할 업체 한 곳을 선택하고 선정사유를 작성합니다.',
    '발주품의 요청을 누른 뒤 통합결재함에서 승인되면 선정업체 기준으로 발주가 확정됩니다.',
])

new_page(doc)
add_heading(doc, '7  기성 및 지급 현황')
screen_frame(doc, '기성 및 지급 현황', '발주별 기성 등록과 실제 지급 결재', ['기성 및 지급 현황', '계약 발주 이력', '준공정산'],
    cards=[('현재 발주금액', '110,200,000 원', '승인된 현재 차수 발주'), ('누적 기성금액', '80,000,000 원', '승인 기성 누계'), ('실제 지급금액', '70,000,000 원', '실제 송금 완료 누계')],
    columns=[('기성 등록', ['발주 선택', '차수와 청구금액 입력', '기성 결재요청']), ('지급 요청', ['승인 후 미지급 발주 선택', '지급결재 요청', '결재 승인 후 실제 지급 반영']), ('잔액 확인', ['미기성금액', '승인 후 미지급', '총 미지급 잔액'])],
    note='현재 승인 차수의 발주만 집계됩니다. 이전 차수 발주와 기성 지급 기록은 보존되며 현재 차수 총액에서는 분리됩니다.')
add_table(doc, ['구분', '뜻', '업무 판단'], [
    ['미기성금액', '현재 발주금액 - 승인 기성누계', '공정이 끝나지 않았거나 기성 승인이 남은 금액'],
    ['승인 후 미지급', '승인 기성누계 - 실제 지급누계', '지급결재 대상 금액'],
    ['총 미지급 잔액', '현재 발주금액 - 실제 지급누계', '향후 집행해야 할 전체 발주 잔액'],
], [1.5, 2.55, 2.8])

new_page(doc)
add_heading(doc, '8  통합결재함 처리')
screen_frame(doc, '통합결재함', '실행예산 변경 차수 발주 기성 정산 결재를 한 화면에서 처리', ['통합결재함', '프로젝트 현황', '변경관리 이력'],
    columns=[('필터', ['전체 상태', '결재대기 승인 반려 승인취소', '전체 유형']), ('결재 목록', ['요청일', '프로젝트와 유형', '제목 요청금액 증감 상태']), ('결재 상세', ['요청 내용과 첨부 확인', '승인 또는 반려', '반려 의견 입력'])],
    note='목록에는 순번과 요청 시간이 분 단위로 표시됩니다. 결재 후에는 관련 프로젝트의 금액과 상태가 자동 갱신됩니다.')
add_heading(doc, '결재 처리 기준', 2)
add_table(doc, ['유형', '승인 전 확인', '승인 후 결과'], [
    ['실행예산', '산출근거, 목표원가한도, 첨부파일', '업체 비교견적 선택 가능'],
    ['계약 실행예산 변경', '증감사유와 변경 차수', '현재 승인 차수 금액으로 집계'],
    ['발주', '선정업체, 견적 대비 차이, 선정사유', '선정업체 발주 확정'],
    ['기성 및 지급', '발주 잔액, 기성 금액, 지급 대상', '기성 또는 실제 지급 누계 갱신'],
    ['준공정산', '미지급 사유, 최종 원가와 이익', '정산완료 처리'],
], [1.55, 3.15, 2.15])

new_page(doc)
add_heading(doc, '9  계약 변경 차수와 이력 관리')
screen_frame(doc, '계약 변경 차수 관리', '계약금액 또는 계약 공종 변경', ['변경관리 이력', '프로젝트 현황', '통합결재함'],
    columns=[('대상 프로젝트', ['P-2026-001', '성가수녀원 리뉴얼공사']), ('새 차수 작성', ['차수 제목', '변경 사유', '새 차수 작성']), ('관리 원칙', ['최초 금액 보존', '승인 증감액 합산', '현재금액 자동 계산'])],
    note='세부실행예산은 세부실행예산 메뉴에서 별도로 작성하고 결재합니다. 계약 변경 차수 화면은 계약금액과 계약 공종 변경만 관리합니다.')
add_bullets(doc, [
    '최초 승인된 금액을 직접 덮어쓰지 않습니다. 새 차수를 작성해 변경 사유와 증감액을 남깁니다.',
    '현재 계약금액, 현재 실행예산, 현재 발주금액은 각각 승인된 증감 이력의 합으로 계산됩니다.',
    '현재 차수가 변경되면 이전 차수의 발주, 기성, 지급은 이력으로 보존되고 현재 집계에서는 분리될 수 있습니다.',
    '프로젝트별 변경 내역이 필요하면 프로젝트를 선택한 뒤 변경관리 이력 메뉴를 엽니다.',
])

new_page(doc)
add_heading(doc, '10  현장 작업일지와 현장 회의록')
add_para(doc, '현장 기록은 별도 결재 없이 프로젝트 이력으로 관리합니다. 공사 진행 중에는 작업일지로 일일 현황을 남기고, 협의 사항은 회의록과 후속조치 항목으로 추적합니다.')
screen_frame(doc, '현장 기록', '현장 작업일지와 현장 회의록', ['현장 작업일지', '현장 회의록', '기성 및 지급 현황'],
    columns=[('작업일지', ['작업일자와 날씨', '투입 인원과 장비', '작업 내용과 특이사항']), ('회의록', ['회의 일시 장소 주제', '참석자와 논의 결정사항', '후속조치 담당자 기한 상태']), ('첨부', ['현장 사진', '도면과 관련문서', '저장 후 파일 첨부'])],
    note='정산완료 프로젝트는 현장 작업일지와 회의록이 조회 전용으로 전환됩니다.')
add_heading(doc, '기록 품질 기준', 2)
add_bullets(doc, [
    '작업일지는 실제 작업일에 작성하고, 진행률이나 이슈가 있으면 구체적인 위치와 내용을 남깁니다.',
    '회의록에는 결정사항과 후속조치를 분리해 기록하고 담당자와 기한을 반드시 지정합니다.',
    '사진, 변경도면, 협의문 등 근거자료는 해당 기록을 저장한 뒤 첨부합니다.',
])

new_page(doc)
add_heading(doc, '11  준공정산과 최종결재')
screen_frame(doc, '준공정산', '계약 실행예산 발주 기성 지급 누계와 최종 손익', ['준공정산', '기성 및 지급 현황', '통합결재함'],
    cards=[('최종 계약금액', '220,000,000 원', '최초 금액과 승인 변경 반영'), ('최종 발주금액', '110,200,000 원', '발주 및 집행 현황'), ('최종 이익률', '49.91%', '최종 계약금액 대비')],
    columns=[('집계 확인', ['계약 및 매출 현황', '실행예산 현황', '발주 기성 지급 현황']), ('직접 정산비 입력', ['미수금', '현장경비', '인건비와 준공 세금']), ('완료 조건', ['미지급 잔액 확인', '사유와 정산 비고', '준공정산 최종 결재요청'])],
    note='미지급 잔액이 남아 있으면 강제 완료 확인과 미지급 처리 사유가 필요합니다. 모든 조건을 충족하면 최종결재를 요청합니다.')
add_table(doc, ['정산 확인 항목', '확인 내용', '결재 전 조치'], [
    ['최종 계약금액', '최초 계약과 승인된 계약 변경 반영 여부', '누락된 계약변경이 있으면 먼저 차수 처리'],
    ['최종 원가', '최종 발주금액과 현장경비 인건비 세금 반영 여부', '직접 정산비를 실제 집행액으로 입력'],
    ['미지급 잔액', '발주 기준 실제 지급 전 남은 금액', '지급 완료 또는 강제 완료 사유 작성'],
    ['최종 이익', '최종 계약금액 - 최종 원가', '예상 이익과 차이가 큰 경우 원인 확인'],
], [1.45, 3.35, 2.05])

new_page(doc)
add_heading(doc, '12  업무 규칙과 자주 묻는 질문')
add_heading(doc, '반드시 지키는 원칙', 2)
add_table(doc, ['원칙', '설명'], [
    ['최초 승인 금액 보존', '승인된 최초 계약 실행예산 발주금액은 직접 수정하지 않습니다. 변경은 별도 차수 또는 변경 이력으로 남깁니다.'],
    ['결재와 실제 지급 분리', '기성 승인과 실제 송금은 다른 단계입니다. 승인된 기성 중 실제 지급이 남으면 승인 후 미지급으로 관리합니다.'],
    ['발주 가능 항목 제한', '업체 비교견적과 발주품의는 승인된 실행예산 항목만 진행할 수 있습니다.'],
    ['현재 차수 기준 집계', '이전 차수의 기록은 보존하되, 현재 프로젝트 통합현황은 현재 승인 차수 기준으로 표시합니다.'],
    ['정산완료 후 조회 전용', '정산완료 프로젝트는 기록 보존을 우선하며 신규 입력과 변경이 제한될 수 있습니다.'],
], [1.65, 5.2])
add_heading(doc, '자주 묻는 질문', 2)
add_table(doc, ['질문', '답변'], [
    ['저장했는데 업체 비교견적에서 항목이 보이지 않습니다.', '세부실행예산을 결재요청하고 승인받았는지 확인합니다. 승인된 항목만 선택할 수 있습니다.'],
    ['발주금액과 기성금액이 왜 다릅니까.', '발주는 약정 금액이고 기성은 공정에 따라 승인된 청구 누계입니다. 지급은 실제 송금 누계이므로 세 값은 서로 다를 수 있습니다.'],
    ['계약금액을 바꿔야 합니다.', '기존 계약을 직접 수정하지 말고 계약 변경안과 변경 차수를 작성해 결재를 받습니다.'],
    ['정산 결재요청 버튼이 비활성화되어 있습니다.', '미지급 잔액, 필수 사유, 강제 완료 확인, 정산비 입력값을 확인한 뒤 다시 시도합니다.'],
], [2.55, 4.3])
add_para(doc, '문서 끝', 9, False, MUTED, 0, WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print(OUT)
